"""Genera la presentación .pptx a partir de un guion en YAML.

Uso:
    .venv/bin/python scripts/build_pptx.py presentacion/diapositivas.yaml [-o presentacion/sesion.pptx]

Tipos de diapositiva admitidos (campo `tipo`):
    portada       titulo, subtitulo
    seccion       titulo, subtitulo?
    contenido     titulo, puntos[], imagen?, fuente?
    imagen        titulo, imagen, pie?, fuente?
    dos_columnas  titulo, izquierda{titulo?, puntos[]}, derecha{titulo?, puntos[]}, fuente?
    tabla         titulo, cabecera[], filas[[]], fuente?
    mensaje       texto, subtexto?          (mensaje clave a pantalla completa)
    referencias   titulo?, referencias[]

Si el YAML incluye `plantilla: ruta/a/plantilla.pptx`, se usan los diseños, colores y tipografía de esa
plantilla (p. ej. el tema «Dividendo» de PowerPoint, convertido con scripts/plantilla_desde_thmx.py)
en lugar del estilo sobrio propio.

Todas admiten `notas` (notas del orador). Un punto puede ser texto o {texto, sub: [..]}.
Las rutas de imagen son relativas a la raíz del repositorio.

Estilo: sobrio y elegante. Blanco y negro con escala de grises y un único acento morado,
usado con moderación (filetes, viñetas, cifras destacadas). Sin bloques de color ni degradados.
El rosa oro es exclusivo de la web: no se usa en la presentación.
"""

import argparse
import sys
from pathlib import Path

import yaml
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

RAIZ = Path(__file__).resolve().parent.parent
ANCHO, ALTO = Inches(13.333), Inches(7.5)
MARGEN = Inches(0.6)

TEMA_POR_DEFECTO = {
    "primario": "111111",    # títulos y texto principal (casi negro)
    "acento": "5B3F8C",      # morado sobrio: único color, usar con moderación
    "texto": "1A1A1A",
    "suave": "707070",       # texto secundario, fuentes, numeración
    "linea": "D4D4D4",       # filetes finos
    "fondo": "FFFFFF",
    "fondo_alt": "F5F5F5",   # gris muy claro (cabeceras de tabla, cajas)
    "oscuro": "111111",      # fondo de las diapositivas de mensaje
    "fuente": "Aptos",
    "fuente_titulo": "Aptos Display",
}


def rgb(hex_color):
    return RGBColor.from_string(hex_color.lstrip("#").upper())


class Constructor:
    def __init__(self, guion):
        self.guion = guion
        self.tema = {**TEMA_POR_DEFECTO, **(guion.get("tema") or {})}
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = ANCHO, ALTO
        self.blank = self.prs.slide_layouts[6]
        self.numero = 0

    # ---------- utilidades ----------
    def _fondo(self, slide, color):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(color)

    def _rect(self, slide, x, y, w, h, color):
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(color)
        shp.line.fill.background()
        return shp

    def _texto(self, slide, x, y, w, h, texto, size=20, bold=False, color=None,
               align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, fuente=None):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for i, linea in enumerate(str(texto).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            run = p.add_run()
            run.text = linea
            f = run.font
            f.size, f.bold, f.italic = Pt(size), bold, italic
            f.name = fuente or self.tema["fuente"]
            f.color.rgb = rgb(color or self.tema["texto"])
        return tb

    def _linea(self, slide, x, y, w, color=None, grosor=Pt(0.75)):
        return self._rect(slide, x, y, w, grosor, color or self.tema["linea"])

    def _titulo(self, slide, titulo):
        self._texto(slide, MARGEN, Inches(0.4), ANCHO - 2 * MARGEN, Inches(0.85),
                    titulo, size=30, color=self.tema["primario"], anchor=MSO_ANCHOR.BOTTOM,
                    fuente=self.tema["fuente_titulo"])
        self._linea(slide, MARGEN, Inches(1.32), ANCHO - 2 * MARGEN)
        self._linea(slide, MARGEN, Inches(1.32) - Pt(0.5), Inches(0.6), self.tema["acento"], Pt(1.75))

    def _pie(self, slide, fuente=None):
        self.numero += 1
        y = ALTO - Inches(0.5)
        if fuente:
            self._texto(slide, MARGEN, y, ANCHO - Inches(2.2), Inches(0.35),
                        f"Fuente: {fuente}", size=10, color=self.tema["suave"])
        self._texto(slide, ANCHO - MARGEN - Inches(0.8), y, Inches(0.8), Inches(0.35),
                    str(self.numero), size=10, color=self.tema["suave"], align=PP_ALIGN.RIGHT)

    def _puntos(self, slide, x, y, w, h, puntos, size=None):
        n = sum(1 + len(p.get("sub", [])) if isinstance(p, dict) else 1 for p in puntos)
        size = size or (24 if n <= 4 else 22 if n <= 6 else 18 if n <= 9 else 16)
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        primero = True

        def add(texto, nivel):
            nonlocal primero
            p = tf.paragraphs[0] if primero else tf.add_paragraph()
            primero = False
            p.space_after = Pt(12 if nivel == 0 else 4)
            viñeta = "—  " if nivel == 0 else "·  "
            r1 = p.add_run()
            r1.text = ("      " * nivel) + viñeta
            r1.font.color.rgb = rgb(self.tema["acento"] if nivel == 0 else self.tema["suave"])
            r1.font.size = Pt(size - 4 * nivel)
            r1.font.name = self.tema["fuente"]
            r2 = p.add_run()
            r2.text = str(texto)
            r2.font.size = Pt(size - 4 * nivel)
            r2.font.name = self.tema["fuente"]
            r2.font.color.rgb = rgb(self.tema["texto"] if nivel == 0 else self.tema["suave"])

        for punto in puntos:
            if isinstance(punto, dict):
                add(punto["texto"], 0)
                for sub in punto.get("sub", []):
                    add(sub, 1)
            else:
                add(punto, 0)
        return tb

    def _imagen(self, slide, ruta, x, y, w, h):
        path = RAIZ / ruta
        if not path.exists():
            print(f"AVISO: imagen no encontrada: {ruta}", file=sys.stderr)
            caja = self._rect(slide, x, y, w, h, self.tema["fondo_alt"])
            caja.text_frame.text = f"[Falta imagen]\n{ruta}"
            return
        with Image.open(path) as im:
            iw, ih = im.size
        escala = min(w / iw, h / ih)
        pw, ph = int(iw * escala), int(ih * escala)
        slide.shapes.add_picture(str(path), x + Emu((w - pw) // 2), y + Emu((h - ph) // 2), pw, ph)

    def _notas(self, slide, datos):
        if datos.get("notas"):
            slide.notes_slide.notes_text_frame.text = str(datos["notas"]).strip()

    # ---------- tipos de diapositiva ----------
    def portada(self, d):
        s = self.prs.slides.add_slide(self.blank)
        self._fondo(s, self.tema["fondo"])
        self._linea(s, MARGEN, Inches(2.0), Inches(0.9), self.tema["acento"], Pt(2))
        self._texto(s, MARGEN, Inches(2.2), ANCHO - 2 * MARGEN, Inches(2.0),
                    d.get("titulo", self.guion.get("titulo", "")), size=48,
                    color=self.tema["primario"], anchor=MSO_ANCHOR.TOP,
                    fuente=self.tema["fuente_titulo"])
        sub = d.get("subtitulo", self.guion.get("subtitulo", ""))
        if sub:
            self._texto(s, MARGEN, Inches(4.2), ANCHO - 2 * MARGEN, Inches(1.0), sub,
                        size=22, color=self.tema["suave"])
        meta = "  ·  ".join(x for x in [self.guion.get("autora"), self.guion.get("fecha")] if x)
        if meta:
            self._linea(s, MARGEN, ALTO - Inches(1.2), ANCHO - 2 * MARGEN)
            self._texto(s, MARGEN, ALTO - Inches(1.05), ANCHO - 2 * MARGEN, Inches(0.5), meta,
                        size=14, color=self.tema["suave"])
        self._notas(s, d)
        self.numero += 1

    def seccion(self, d):
        s = self.prs.slides.add_slide(self.blank)
        self._fondo(s, self.tema["fondo"])
        self.seccion_n = getattr(self, "seccion_n", 0) + 1
        self._texto(s, MARGEN, Inches(2.2), Inches(3), Inches(0.5), f"{self.seccion_n:02d}",
                    size=16, color=self.tema["acento"])
        self._texto(s, MARGEN, Inches(2.7), ANCHO - 2 * MARGEN, Inches(1.4),
                    d["titulo"], size=40, color=self.tema["primario"],
                    anchor=MSO_ANCHOR.TOP, fuente=self.tema["fuente_titulo"])
        if d.get("subtitulo"):
            self._texto(s, MARGEN, Inches(4.1), ANCHO - 2 * MARGEN, Inches(0.8),
                        d["subtitulo"], size=20, color=self.tema["suave"])
        self._notas(s, d)
        self.numero += 1

    def contenido(self, d):
        s = self.prs.slides.add_slide(self.blank)
        self._fondo(s, self.tema["fondo"])
        self._titulo(s, d["titulo"])
        top, alto = Inches(1.5), ALTO - Inches(2.2)
        if d.get("imagen"):
            ancho_txt = (ANCHO - 2 * MARGEN) * 0.52
            self._puntos(s, MARGEN, top, ancho_txt, alto, d.get("puntos", []))
            xi = MARGEN + ancho_txt + Inches(0.3)
            self._imagen(s, d["imagen"], xi, top, ANCHO - MARGEN - xi, alto)
        else:
            self._puntos(s, MARGEN, top, ANCHO - 2 * MARGEN, alto, d.get("puntos", []))
        self._pie(s, d.get("fuente"))
        self._notas(s, d)

    def imagen(self, d):
        s = self.prs.slides.add_slide(self.blank)
        self._fondo(s, self.tema["fondo"])
        self._titulo(s, d["titulo"])
        alto = ALTO - Inches(2.3) - (Inches(0.5) if d.get("pie") else 0)
        self._imagen(s, d["imagen"], MARGEN, Inches(1.45), ANCHO - 2 * MARGEN, alto)
        if d.get("pie"):
            self._texto(s, MARGEN, Inches(1.45) + alto + Inches(0.05), ANCHO - 2 * MARGEN,
                        Inches(0.5), d["pie"], size=16, color=self.tema["suave"],
                        align=PP_ALIGN.CENTER)
        self._pie(s, d.get("fuente"))
        self._notas(s, d)

    def dos_columnas(self, d):
        s = self.prs.slides.add_slide(self.blank)
        self._fondo(s, self.tema["fondo"])
        self._titulo(s, d["titulo"])
        ancho = (ANCHO - 2 * MARGEN - Inches(0.4)) / 2
        for i, lado in enumerate(["izquierda", "derecha"]):
            col = d.get(lado, {})
            x = MARGEN + i * (ancho + Inches(0.4))
            y = Inches(1.5)
            if col.get("titulo"):
                self._texto(s, x, y, ancho, Inches(0.55), col["titulo"], size=20,
                            bold=True, color=self.tema["primario"], anchor=MSO_ANCHOR.MIDDLE)
                self._linea(s, x, y + Inches(0.6), ancho)
                y += Inches(0.8)
            self._puntos(s, x, y, ancho, ALTO - y - Inches(0.7), col.get("puntos", []), size=18)
        self._pie(s, d.get("fuente"))
        self._notas(s, d)

    def tabla(self, d):
        s = self.prs.slides.add_slide(self.blank)
        self._fondo(s, self.tema["fondo"])
        self._titulo(s, d["titulo"])
        cab, filas = d["cabecera"], d["filas"]
        alto_fila = Inches(0.55)
        t = s.shapes.add_table(len(filas) + 1, len(cab), MARGEN, Inches(1.6),
                               ANCHO - 2 * MARGEN, alto_fila * (len(filas) + 1)).table
        size = 16 if len(filas) <= 6 else 13
        for r, fila in enumerate([cab] + filas):
            for c, valor in enumerate(fila):
                celda = t.cell(r, c)
                celda.text = str(valor)
                for p in celda.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(size)
                        run.font.name = self.tema["fuente"]
                        run.font.bold = r == 0
                        run.font.color.rgb = rgb(self.tema["primario"] if r == 0 else self.tema["texto"])
                celda.fill.solid()
                celda.fill.fore_color.rgb = rgb(self.tema["fondo_alt"] if r == 0 else self.tema["fondo"])
        # Sin estilo de tabla de Office: filetes finos entre filas
        t._tbl.tblPr.set("bandRow", "0")
        t._tbl.tblPr.set("firstRow", "0")
        for r in range(1, len(filas) + 1):
            self._linea(s, MARGEN, Inches(1.6) + alto_fila * r, ANCHO - 2 * MARGEN)
        self._pie(s, d.get("fuente"))
        self._notas(s, d)

    def mensaje(self, d):
        s = self.prs.slides.add_slide(self.blank)
        self._fondo(s, self.tema["oscuro"])
        self._linea(s, (ANCHO - Inches(0.9)) // 2, Inches(1.9), Inches(0.9), "9C85C4", Pt(2))
        self._texto(s, Inches(1.4), Inches(2.1), ANCHO - Inches(2.8), Inches(2.6), d["texto"],
                    size=36, color="FFFFFF", align=PP_ALIGN.CENTER,
                    anchor=MSO_ANCHOR.MIDDLE, fuente=self.tema["fuente_titulo"])
        if d.get("subtexto"):
            self._texto(s, Inches(1.4), Inches(4.8), ANCHO - Inches(2.8), Inches(1.2),
                        d["subtexto"], size=18, color="BDBDBD", align=PP_ALIGN.CENTER)
        self._notas(s, d)
        self.numero += 1

    def referencias(self, d):
        s = self.prs.slides.add_slide(self.blank)
        self._fondo(s, self.tema["fondo"])
        self._titulo(s, d.get("titulo", "Bibliografía"))
        refs = d["referencias"]
        size = 14 if len(refs) <= 8 else 12 if len(refs) <= 12 else 10
        texto = "\n".join(f"{i}. {r}" for i, r in enumerate(refs, 1))
        self._texto(s, MARGEN, Inches(1.5), ANCHO - 2 * MARGEN, ALTO - Inches(2.2), texto,
                    size=size)
        self._pie(s)
        self._notas(s, d)

    def construir(self, salida):
        for i, d in enumerate(self.guion["diapositivas"], 1):
            tipo = d.get("tipo", "contenido")
            metodo = getattr(self, tipo, None)
            if metodo is None or tipo.startswith("_") or tipo == "construir":
                sys.exit(f"Diapositiva {i}: tipo desconocido '{tipo}'")
            metodo(d)
        salida.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(salida)
        print(f"Generada {salida} ({len(self.prs.slides)} diapositivas)")


class ConstructorPlantilla:
    """Construye la presentación sobre los diseños de una plantilla .pptx (clave `plantilla` del YAML),
    por ejemplo el tema «Dividendo» de PowerPoint convertido con scripts/plantilla_desde_thmx.py.
    Colores y tipografía los pone el tema: aquí solo se colocan contenidos en sus marcadores."""

    TITULO, CONTENIDO, SECCION, DOS, COMPARACION, SOLO_TITULO = 0, 1, 2, 3, 4, 5
    GRIS = "7F7F7F"

    def __init__(self, guion):
        self.guion = guion
        self.prs = Presentation(RAIZ / guion["plantilla"])
        self.numero = 0

    # ---------- utilidades ----------
    def _nueva(self, diseno):
        self.numero += 1
        return self.prs.slides.add_slide(self.prs.slide_layouts[diseno])

    @staticmethod
    def _ph(slide, idx):
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == idx:
                return ph
        return None

    @staticmethod
    def _quitar(ph):
        ph._element.getparent().remove(ph._element)

    def _rellenar(self, ph, puntos, size=20):
        tf = ph.text_frame
        tf.word_wrap = True
        primero = True

        def add(texto, nivel):
            nonlocal primero
            p = tf.paragraphs[0] if primero else tf.add_paragraph()
            primero = False
            p.text = str(texto)
            p.level = nivel
            for r in p.runs:
                r.font.size = Pt(size if nivel == 0 else size - 3)

        for punto in puntos:
            if isinstance(punto, dict):
                add(punto["texto"], 0)
                for sub in punto.get("sub", []):
                    add(sub, 1)
            else:
                add(punto, 0)

    def _texto(self, slide, x, y, w, h, texto, size=11, color=None, align=None):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = texto
        if align:
            p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.color.rgb = rgb(color or self.GRIS)
        return tb

    def _pie(self, slide, fuente=None):
        if fuente:
            self._texto(slide, Inches(0.64), Inches(6.51), Inches(10.6), Inches(0.4), fuente, size=11)
        self._texto(slide, Inches(11.55), Inches(6.51), Inches(1.15), Inches(0.4), str(self.numero),
                    size=11, align=PP_ALIGN.RIGHT)

    def _imagen(self, slide, ruta, x, y, w, h):
        path = RAIZ / ruta
        if not path.exists():
            print(f"AVISO: imagen no encontrada: {ruta}", file=sys.stderr)
            self._texto(slide, x, y, w, h, f"[Falta imagen] {ruta}", size=14)
            return
        with Image.open(path) as im:
            iw, ih = im.size
        escala = min(w / iw, h / ih)
        pw, ph = int(iw * escala), int(ih * escala)
        slide.shapes.add_picture(str(path), x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)

    def _notas(self, slide, d):
        if d.get("notas"):
            slide.notes_slide.notes_text_frame.text = str(d["notas"]).strip()

    def _titulo(self, slide, texto):
        slide.shapes.title.text = texto

    # ---------- tipos de diapositiva ----------
    def portada(self, d):
        s = self._nueva(self.TITULO)
        self._titulo(s, self.guion.get("titulo", ""))
        sub = self._ph(s, 1)
        lineas = [self.guion.get("subtitulo", "")]
        meta = " · ".join(x for x in (self.guion.get("autora"), self.guion.get("fecha")) if x)
        if meta:
            lineas.append(meta)
        sub.text_frame.text = lineas[0]
        if meta:  # autora y fecha sobre la banda de color del tema, en blanco
            self._texto(s, Inches(0.64), Inches(5.55), Inches(12.0), Inches(0.5), meta, size=18,
                        color="FFFFFF")
        self._notas(s, d)

    def seccion(self, d):
        s = self._nueva(self.SECCION)
        self._titulo(s, d["titulo"])
        cuerpo = self._ph(s, 1)
        if d.get("subtitulo"):
            cuerpo.text_frame.text = d["subtitulo"]
        else:
            self._quitar(cuerpo)
        self._notas(s, d)

    def mensaje(self, d):
        s = self._nueva(self.SECCION)
        self._titulo(s, d["texto"])
        cuerpo = self._ph(s, 1)
        if d.get("subtexto"):
            cuerpo.text_frame.text = d["subtexto"]
        else:
            self._quitar(cuerpo)
        self._notas(s, d)

    def contenido(self, d):
        puntos = d.get("puntos", [])
        n = sum(1 + len(p.get("sub", [])) if isinstance(p, dict) else 1 for p in puntos)
        if d.get("imagen"):
            s = self._nueva(self.DOS)
            self._titulo(s, d["titulo"])
            self._rellenar(self._ph(s, 1), puntos, size=18 if n <= 6 else 16)
            der = self._ph(s, 2)
            x, y, w, h = der.left, der.top, der.width, der.height
            self._quitar(der)
            self._imagen(s, d["imagen"], x, y - Inches(0.25), w, h + Inches(0.3))
        else:
            s = self._nueva(self.CONTENIDO)
            self._titulo(s, d["titulo"])
            self._rellenar(self._ph(s, 1), puntos, size=22 if n <= 5 else 20 if n <= 7 else 18)
        self._pie(s, d.get("fuente"))
        self._notas(s, d)

    def imagen(self, d):
        s = self._nueva(self.SOLO_TITULO)
        self._titulo(s, d["titulo"])
        alto = Inches(4.3) - (Inches(0.4) if d.get("pie") else 0)
        self._imagen(s, d["imagen"], Inches(0.64), Inches(2.05), Inches(12.06), alto)
        if d.get("pie"):
            self._texto(s, Inches(0.64), Inches(2.05) + alto + Inches(0.02), Inches(12.06), Inches(0.4),
                        d["pie"], size=16, align=PP_ALIGN.CENTER)
        self._pie(s, d.get("fuente"))
        self._notas(s, d)

    def dos_columnas(self, d):
        s = self._nueva(self.COMPARACION)
        self._titulo(s, d["titulo"])
        for idx_t, idx_c, col in ((1, 2, d.get("izquierda", {})), (3, 4, d.get("derecha", {}))):
            t = self._ph(s, idx_t)
            if col.get("titulo"):
                t.text_frame.text = col["titulo"]
            else:
                self._quitar(t)
            self._rellenar(self._ph(s, idx_c), col.get("puntos", []), size=18)
        self._pie(s, d.get("fuente"))
        self._notas(s, d)

    def tabla(self, d):
        s = self._nueva(self.SOLO_TITULO)
        self._titulo(s, d["titulo"])
        cab, filas = d["cabecera"], d["filas"]
        alto_fila = Inches(0.62)
        t = s.shapes.add_table(len(filas) + 1, len(cab), Inches(0.64), Inches(2.15),
                               Inches(12.06), alto_fila * (len(filas) + 1)).table
        for r, fila in enumerate([cab] + filas):
            for c, valor in enumerate(fila):
                celda = t.cell(r, c)
                celda.text = str(valor)
                celda.vertical_anchor = MSO_ANCHOR.MIDDLE
                for p in celda.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(16 if r else 17)
                        run.font.bold = r == 0
        self._pie(s, d.get("fuente"))
        self._notas(s, d)

    def referencias(self, d):
        s = self._nueva(self.CONTENIDO)
        self._titulo(s, d.get("titulo", "Bibliografía"))
        self._rellenar(self._ph(s, 1), d.get("referencias", []), size=14)
        self._notas(s, d)

    def construir(self, salida):
        for i, d in enumerate(self.guion["diapositivas"], 1):
            tipo = d.get("tipo", "contenido")
            metodo = getattr(self, tipo, None)
            if metodo is None or tipo.startswith("_") or tipo == "construir":
                sys.exit(f"Diapositiva {i}: tipo desconocido '{tipo}'")
            metodo(d)
        salida.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(salida)
        print(f"Generada {salida} ({len(self.prs.slides)} diapositivas, plantilla {self.guion['plantilla']})")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("guion", type=Path)
    ap.add_argument("-o", "--salida", type=Path)
    args = ap.parse_args()
    guion = yaml.safe_load(args.guion.read_text(encoding="utf-8"))
    salida = (args.salida or RAIZ / "presentacion" / guion.get("archivo", "sesion-clinica.pptx")).resolve()
    (ConstructorPlantilla if guion.get("plantilla") else Constructor)(guion).construir(salida)


if __name__ == "__main__":
    main()
